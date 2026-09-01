from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BLACK = RGBColor(0, 0, 0)


def set_run_black(run, font_size: int | None = None, bold: bool | None = None):
    run.font.color.rgb = BLACK
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if font_size is not None:
        run.font.size = Pt(font_size)
    if bold is not None:
        run.bold = bold


def set_style(style, font_size: int, bold: bool = False):
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(font_size)
    style.font.color.rgb = BLACK
    style.font.bold = bold


def remove_style_borders(style):
    properties = style._element.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is not None:
        properties.remove(borders)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_black(run, 9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)
    run = paragraph.add_run(" 页")
    set_run_black(run, 9)


def configure_document(document: Document):
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    styles = document.styles
    set_style(styles["Normal"], 10.5)
    set_style(styles["Title"], 20, True)
    remove_style_borders(styles["Title"])
    set_style(styles["Heading 1"], 15, True)
    set_style(styles["Heading 2"], 12.5, True)
    set_style(styles["Heading 3"], 11, True)

    for style_name in ("List Bullet", "List Number"):
        if style_name in styles:
            set_style(styles[style_name], 10.5)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def add_text(document: Document, text: str, style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_run_black(run)
    paragraph.paragraph_format.space_after = Pt(5)
    return paragraph


def add_markdown_to_docx(markdown_path: Path, docx_path: Path):
    document = Document()
    configure_document(document)
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    first_heading = True

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        if stripped.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(stripped[2:].strip())
            set_run_black(run, 20, True)
            paragraph.paragraph_format.space_after = Pt(14)
            first_heading = False
            continue

        heading = re.match(r"^(#{2,3})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1)) - 1
            paragraph = document.add_paragraph(style=f"Heading {level}")
            run = paragraph.add_run(heading.group(2).strip())
            set_run_black(run, 15 if level == 1 else 12.5 if level == 2 else 11, True)
            paragraph.paragraph_format.space_before = Pt(9)
            paragraph.paragraph_format.space_after = Pt(4)
            continue

        if re.match(r"^-\s+", stripped):
            text = re.sub(r"^-\s+", "", stripped)
            paragraph = document.add_paragraph(style="List Bullet")
            run = paragraph.add_run(text)
            set_run_black(run)
            paragraph.paragraph_format.space_after = Pt(2)
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            run = paragraph.add_run(numbered.group(1))
            set_run_black(run)
            paragraph.paragraph_format.space_after = Pt(2)
            continue

        add_text(document, stripped)

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            set_run_black(run)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(docx_path)


def main():
    jobs = [
        (
            ROOT / "docs" / "八字紫微排盘产品PRD.md",
            ROOT / "docs" / "八字紫微排盘产品PRD-微信小程序版.docx",
        ),
        (
            ROOT / "discuss" / "八字紫微排盘Phase计划.md",
            ROOT / "discuss" / "八字紫微排盘Phase计划-微信小程序版.docx",
        ),
    ]
    for markdown_path, docx_path in jobs:
        add_markdown_to_docx(markdown_path, docx_path)
        print(docx_path)


if __name__ == "__main__":
    main()
